"""
tools/make_demo_deck.py - the deck you read off while demoing.

    python tools/make_demo_deck.py   ->  docs/NEXUS_Demo.pptx

This is NOT the pitch deck. It has one job: be readable from a laptop screen
while you are talking and clicking, so nothing has to be memorised.

Design rules, which are different from a normal deck on purpose:
  - One time block per slide, in clock order, with the clock time huge in the
    corner. You should be able to glance at it and know if you are behind.
  - What to CLICK in amber, what to SAY in white quotes. Never mixed.
  - Big type. This is read at a glance from a metre away, mid-sentence.
  - The criterion each block earns is printed on it, because the scoring
    sheet has six lines and you should be able to see you have covered them.
  - A picture section at the end: what each moment actually looks like, so a
    judge who missed a click can still see what happened.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
SHOTS = os.path.join(DOCS, 'screenshots', 'demo')
OUT = os.path.join(DOCS, 'NEXUS_Demo.pptx')

BG = '#0E1116'
PANEL = '#171C24'
TEXT = '#E8ECF2'
DIM = '#8A94A6'
AMBER = '#F5A524'
GREEN = '#2ECC71'
RED = '#E5484D'
BLUE = '#4C9AFF'

FONT = 'Segoe UI'
MONO = 'Consolas'
SW, SH = Inches(13.333), Inches(7.5)


def rgb(h):
    return RGBColor(int(h[1:3], 16), int(h[3:5], 16), int(h[5:7], 16))


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BG)
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def text(slide, x, y, w, h, body, size=20, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, spacing=1.2, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(body.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = rgb(color)
    return tb


def bar(slide, y, x=0.7, w=11.9, colour=AMBER, h=0.05):
    r = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(colour)
    r.line.fill.background()
    r.shadow.inherit = False


def clock(slide, when):
    """The time this block starts, big, top right. Glanceable."""
    text(slide, 9.6, 0.35, 3.1, 1.0, when, size=44, bold=True,
         color=AMBER, font=MONO, align=PP_ALIGN.RIGHT)


def earns(slide, label, colour=BLUE):
    text(slide, 0.75, 6.62, 11.8, 0.5, 'EARNS:  ' + label.upper(),
         size=13, bold=True, color=colour, font=MONO)


# =============================================================================
# THE RUN SHEET - one slide per block, in clock order
# =============================================================================
# (time, title, [click lines], say, criterion, colour)
BLOCKS = [
    ('0:00', 'Point at both windows',
     ['Nothing to click. Just point.'],
     '"Two screens, two jobs.\n\n'
     'Left is THE STREET — the digital twin. Four junctions, real traffic '
     'engineering numbers. Five vehicle types, because Indian traffic is not '
     'all cars.\n\n'
     'Right is THE CONTROL ROOM — what an operator sees. It has ZERO physics '
     'in it. Everything it knows arrives as a message."',
     'UI/UX & User Experience', BLUE),

    ('0:45', 'The story',
     ['Still nothing to click.'],
     '"In India most families do not wait for an ambulance. They drive. '
     'That happened in my family, with my grandmother.\n\n'
     'The car was not slow. The time was lost standing still, at one red '
     'light after another.\n\n'
     'It hit a wall that was not technical. It was trust. Give one vehicle a '
     'green and somebody else pays for it — and no traffic department will '
     'discuss it until you can say what it costs.\n\n'
     'Nobody could answer that. So I built the place to try it."',
     'Remarks / problem-solution fit', GREEN),

    ('1:30', 'Change a timing. Prove it is better',
     ['PLAN tab',
      'Point at the SEED number',
      'N–S GREEN -> 30      E–W GREEN -> 8',
      'TEST THIS PLAN — 100s',
      'SUGGEST A PLAN (WEBSTER 1958)',
      'TEST THIS PLAN — 100s   again'],
     '"Every test replays identical traffic. Same seed. So the difference is '
     'the plan, not luck.\n\n'
     'That second one is Webster\'s minimum-delay cycle, from 1958 — what most '
     'Indian junctions actually run.\n\n'
     'Most dashboards tell you what IS happening. This one tells you what '
     'WOULD happen if you changed something, and what it would cost."',
     'Functionality & Completeness   ←   NEVER CUT THIS', AMBER),

    ('3:00', 'The bad day — one system, two screens',
     ['OPERATE tab',
      'ACCIDENT   (point at the red X)',
      'FLOOD',
      'NOW TURN TO THE CONTROL ROOM'],
     '"A collision at J2 — that approach cannot discharge. And a flooded road: '
     'heavy vehicles barred by weight restriction.\n\n'
     'The operator\'s screen did not need telling. THE NETWORK IS DEGRADED, two '
     'conditions — and it says what each one MEANS, not just what it is.\n\n'
     'That is the difference between a dashboard and a control room."',
     'Functionality & Completeness  +  UI/UX', AMBER),

    ('4:30', 'Two ambulances — decide, then veto',
     ['ALL CLEAR   then   TWO AMBULANCES',
      'Point at the two white ambulances',
      'TURN TO CONTROL ROOM (countdown)',
      'Click DO NOT PROCEED',
      'Fire TWO AMBULANCES again — let it run',
      'Watch a tag flip HELD -> PRIORITY'],
     '"Two ambulances, one junction, from adjacent sides. Only one phase can '
     'be green. There is no version where nobody waits.\n\n'
     'It shows its work in two columns, deliberately not merged. THE CASE — '
     'severity, people on board. THE NETWORK — what the trained model values '
     'that junction at. The model is a modest adjustment, not an equal term. '
     'I will not let it outvote a critical patient.\n\n'
     'Default is to ACT. A system that waits for a human to click before an '
     'ambulance gets a green has moved the delay, not removed it."',
     'Technical Execution  +  UI/UX   ←   YOUR CENTREPIECE', RED),

    ('6:30', 'The model, and the model that checks it',
     ['MODEL tab',
      'TRAINED MODEL',
      'LLM AGENT   (point at the three layers)'],
     '"Reinforcement learning, trained against THIS simulator running headless '
     '— same physics, no second implementation to drift.\n\n'
     'Fixed plan 36.9s.   Max-Pressure 33.5s.   This model 32.6s.\n\n'
     'And the part that did NOT work: a standard DQN never beat the fixed '
     'plan. Monte-Carlo returns did.\n\n'
     'One model proposes an action and a reason. A SECOND call reviews it and '
     'can overrule it. Under both, a local rule whenever the API is down — and '
     'the screen says which layer decided."',
     'Technical Execution', BLUE),

    ('7:30', 'Quantum-ready, and a real corridor',
     ['SOLVE THE NETWORK AS A QUBO',
      '(cut this slide first if behind)'],
     '"Every controller so far decides one junction at a time. This decides '
     'all four at once, written as a QUBO — the exact form a quantum annealer '
     'takes as input.\n\n'
     'Solved here CLASSICALLY, on this laptop. Nothing about it is quantum. '
     'But four junctions is only 256 states, so I brute-forced the optimum and '
     'checked the annealer found it. It did. 256 states here — 2^100 for fifty '
     'junctions. That is why the formulation matters.\n\n'
     'Then re-run in SUMO on 1.71 km of real Vijayawada, 19 junctions: '
     'ambulance 369s to 177s, waiting 160s to zero, everyone else 1.5% BETTER."',
     'Feasibility & Scalability', GREEN),

    ('8:30', '3D',
     ['Click 3D in the top bar',
      'Orbit once with the mouse',
      'Click back to 2D'],
     '"Same simulation, two views. Nothing in the physics knows which one you '
     'are looking at — only the renderer draws. That is why the 3D view was an '
     'afternoon and touched no traffic code.\n\n'
     '2D is the one you make decisions in. 3D is the one you show a mayor."',
     'UI/UX & User Experience', BLUE),

    ('9:00', 'Close — say it before they do',
     ['Nothing to click.'],
     '"Adaptive signal control is not new and I will not pretend it is. '
     'Bengaluru runs B-ATCS on ~165 junctions, ~33% at Hudson Circle. '
     'Pittsburgh\'s Surtrac reports 25.\n\n'
     'So my contribution is narrower: you can PREVIEW a change before deploying '
     'it, every decision states its REASON in numbers you can check, and '
     'priority has a PUBLISHED PRICE.\n\n'
     'One thing I did not expect: OpenStreetMap tags exactly TWO traffic '
     'signals across 516 km of Vijayawada road.\n\n'
     'Before a city changes a signal, it should be able to try the change and '
     'see who pays for it. Ask me anything — including what does not work."',
     'Remarks', GREEN),
]


# =============================================================================
# THE PICTURES - what each moment actually looks like
# =============================================================================
# (file, headline, what it shows)
PICTURES = [
    ('01-street-normal.png',
     'The street, running normally',
     'Four junctions, five vehicle types, live delay and queue counts in the HUD. '
     'Every number here is measured from vehicle state, not estimated.'),

    ('02-street-two-plans-tested.png',
     'Change the timing — and find out if you were right',
     'Two plans tested on IDENTICAL traffic (same seed). The results table keeps both '
     'and marks the better one. This is the deliverable: a planner can try a change '
     'before it goes on a real road.'),

    ('03-street-bad-day.png',
     'The bad day: a collision and a flooded road',
     'The red X marks the approach that cannot discharge — watch the queue build '
     'backwards into it. The blue band is the flooded road: speeds drop and heavy '
     'vehicles are barred.'),

    ('04-room-degraded.png',
     'The control room noticed by itself',
     'THE NETWORK IS DEGRADED. Each condition says what it MEANS for the operator, '
     'not just what it is — a priority corridor down a flooded road will not run at '
     'the speed the estimate assumes.'),

    ('05-street-two-ambulances.png',
     'Two ambulances, one junction, adjacent approaches',
     'Both are real vehicles from the moment the calls come in, arriving within a '
     'second of each other. Only one phase can be green.'),

    ('06-room-plan-and-veto.png',
     'The system decided, showed its work, and can be stopped',
     'THE CASE (severity, people on board) and THE NETWORK (what the trained model '
     'values that junction at) are kept in separate columns. It runs on its own in '
     '8 seconds — the operator vetoes, rather than decides.'),

    ('07-street-corridor-running.png',
     'One served, one held — and the price is on screen',
     'The corridor is running. COST OF PRIORITY shows seconds saved for the ambulance '
     'AND vehicle-seconds paid by everybody else. No deployed system publishes that '
     'second number.'),

    ('08-street-model-and-qubo.png',
     'The trained model, the language-model stack, and the QUBO',
     'Held-out benchmark against two baselines, the three-layer model stack with its '
     'fallback, and the whole network solved as a QUBO — classically, and checked '
     'against brute force.'),

    ('09-street-3d.png',
     'The same simulation, in 3D',
     'Same state, different renderer. Ambulances carry PRIORITY or HELD tags, blocked '
     'approaches show as red bars, flooded roads are tinted. Nothing below the '
     'renderer changed to make this exist.'),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # ------------------------------------------------------------- title
    s = blank(prs)
    bar(s, 2.1, x=0.75, w=2.2)
    text(s, 0.75, 2.45, 11.9, 1.2, 'DEMO RUN SHEET', size=54, bold=True)
    text(s, 0.75, 3.65, 11.9, 0.8, 'Read this. Do not improvise.',
         size=26, color=AMBER, bold=True)
    text(s, 0.75, 4.6, 11.9, 1.4,
         'Amber = what to click.   White quotes = what to say.\n'
         'The clock in the corner is when that block STARTS.',
         size=17, color=DIM)
    text(s, 0.75, 6.3, 11.9, 0.6,
         'TEAM NEXUS   ·   BEFORE YOU START: run run.bat, open both windows side by side',
         size=13, color=DIM, font=MONO, bold=True)

    # -------------------------------------------------------- the blocks
    for when, title, clicks, say, crit, colour in BLOCKS:
        s = blank(prs)
        clock(s, when)
        bar(s, 1.35, x=0.75, w=8.6, colour=colour)
        text(s, 0.75, 0.45, 8.7, 0.9, title, size=30, bold=True)

        y = 1.65
        text(s, 0.75, y, 5.4, 0.4, 'CLICK', size=13, bold=True, color=AMBER, font=MONO)
        y += 0.45
        for c in clicks:
            text(s, 0.75, y, 5.6, 0.5, '›  ' + c, size=15, color=AMBER, font=MONO)
            y += 0.5

        text(s, 6.6, 1.65, 6.0, 0.4, 'SAY', size=13, bold=True, color=DIM, font=MONO)
        text(s, 6.6, 2.05, 6.05, 4.5, say, size=14, spacing=1.26)

        earns(s, crit, colour)

    # ------------------------------------------------------ picture divider
    s = blank(prs)
    bar(s, 2.6, x=0.75, w=2.2)
    text(s, 0.75, 2.95, 11.9, 1.2, 'WHAT IT LOOKS LIKE', size=48, bold=True)
    text(s, 0.75, 4.2, 11.9, 1.0,
         'The same run, photographed. If a judge missed a click, this is what happened.',
         size=20, color=DIM)

    # ------------------------------------------------------- the pictures
    for fname, headline, caption in PICTURES:
        p = os.path.join(SHOTS, fname)
        if not os.path.exists(p):
            print('  missing, skipped: ' + fname)
            continue
        s = blank(prs)
        text(s, 0.75, 0.4, 11.9, 0.7, headline, size=26, bold=True)
        bar(s, 1.12, x=0.75, w=1.6)

        # The picture gets the room; the caption sits under it.
        s.shapes.add_picture(p, Inches(0.75), Inches(1.45), height=Inches(4.55))
        text(s, 0.75, 6.15, 11.9, 1.0, caption, size=14.5, color=DIM, spacing=1.3)

    # ------------------------------------------------------------ the end
    s = blank(prs)
    bar(s, 2.3, x=0.75, w=2.2)
    text(s, 0.75, 2.65, 11.9, 2.2,
         'Before a city changes a signal,\nit should be able to try the change\n'
         'and see who pays for it.',
         size=34, bold=True, spacing=1.25)
    text(s, 0.75, 5.4, 11.9, 0.7, 'Ask me anything — including what does not work.',
         size=22, bold=True, color=AMBER)

    prs.save(OUT)
    return OUT


if __name__ == '__main__':
    print('wrote ' + build())
